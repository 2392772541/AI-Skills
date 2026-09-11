#!/usr/bin/env python3
"""Validate 歌曲目录.json and rebuild the portable segment-first music library."""

from __future__ import annotations

import argparse
import array
import csv
import json
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path, PurePosixPath


REQUIRED = (
    "file_name",
    "type",
    "language",
    "core_emotions",
    "fine_emotions",
    "intensity",
    "atmosphere",
    "uses",
    "relative_path",
)
ARRAY_FIELDS = ("core_emotions", "fine_emotions", "atmosphere", "uses")
HEADERS = (
    "文件名",
    "类型",
    "语言",
    "核心情绪",
    "细分情绪",
    "情绪强度",
    "氛围关键词",
    "适用场景",
    "来源",
    "授权状态",
    "时长",
    "相对路径",
    "备注",
)
SEGMENT_HEADERS = (
    "歌曲",
    "片段名称",
    "开始",
    "结束",
    "片段时长",
    "核心情绪",
    "细分情绪",
    "强度",
    "用途",
    "适合画面",
    "口播适配",
    "人声状态",
    "标注来源",
)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, required=True, help="素材库根目录")
    parser.add_argument("--catalog", default="歌曲目录.json", help="相对根目录的目录文件")
    parser.add_argument("--check-only", action="store_true", help="只校验，不生成文件")
    return parser.parse_args()


def atomic_write_text(path: Path, text: str, encoding: str = "utf-8") -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temp_name = tempfile.mkstemp(prefix=f".{path.name}.", suffix=".tmp", dir=path.parent)
    try:
        with os.fdopen(fd, "w", encoding=encoding, newline="") as handle:
            handle.write(text)
        os.replace(temp_name, path)
    except Exception:
        try:
            os.unlink(temp_name)
        except FileNotFoundError:
            pass
        raise


def parse_duration(value: str) -> float:
    try:
        parts = [float(part) for part in value.split(":")]
    except (AttributeError, ValueError):
        return 0.0
    if len(parts) == 2:
        return parts[0] * 60 + parts[1]
    if len(parts) == 3:
        return parts[0] * 3600 + parts[1] * 60 + parts[2]
    return 0.0


def seconds_label(seconds: float) -> str:
    seconds = max(0, int(round(seconds)))
    return f"{seconds // 60:02d}:{seconds % 60:02d}"


def load_and_validate(root: Path, catalog_path: Path) -> tuple[dict, list[str]]:
    errors: list[str] = []
    try:
        data = json.loads(catalog_path.read_text(encoding="utf-8"))
    except FileNotFoundError:
        return {}, [f"目录文件不存在: {catalog_path}"]
    except json.JSONDecodeError as exc:
        return {}, [f"JSON格式错误: {exc}"]

    if not isinstance(data, dict) or not isinstance(data.get("tracks"), list):
        return data, ["根对象必须包含 tracks 数组"]

    seen_names: set[str] = set()
    seen_paths: set[str] = set()
    for index, track in enumerate(data["tracks"], start=1):
        prefix = f"tracks[{index}]"
        if not isinstance(track, dict):
            errors.append(f"{prefix}: 必须是对象")
            continue
        for field in REQUIRED:
            if field not in track or track[field] in (None, "", []):
                errors.append(f"{prefix}: 缺少必填字段 {field}")
        for field in ARRAY_FIELDS:
            value = track.get(field)
            if not isinstance(value, list) or not value or any(not isinstance(v, str) or not v.strip() for v in value):
                errors.append(f"{prefix}.{field}: 必须是非空字符串数组")
            elif len(value) != len(set(value)):
                errors.append(f"{prefix}.{field}: 包含重复标签")

        name = track.get("file_name")
        rel = track.get("relative_path")
        if isinstance(name, str):
            if name in seen_names:
                errors.append(f"{prefix}: 文件名重复 {name}")
            seen_names.add(name)
        if isinstance(rel, str):
            if "\\" in rel:
                errors.append(f"{prefix}: relative_path 必须使用 / 分隔")
            pure = PurePosixPath(rel)
            if pure.is_absolute() or ".." in pure.parts:
                errors.append(f"{prefix}: relative_path 必须是安全相对路径")
            if rel in seen_paths:
                errors.append(f"{prefix}: 路径重复 {rel}")
            seen_paths.add(rel)
            if isinstance(name, str) and pure.name != name:
                errors.append(f"{prefix}: file_name 与路径末尾不一致")
            if not root.joinpath(*pure.parts).is_file():
                errors.append(f"{prefix}: 音频文件不存在 {rel}")

        segments = track.get("segments")
        if segments is not None:
            if not isinstance(segments, list) or not segments:
                errors.append(f"{prefix}.segments: 必须是非空数组")
            else:
                last_end = 0.0
                for segment_index, segment in enumerate(segments, start=1):
                    segment_prefix = f"{prefix}.segments[{segment_index}]"
                    if not isinstance(segment, dict):
                        errors.append(f"{segment_prefix}: 必须是对象")
                        continue
                    for field in ("label", "start", "end", "uses", "voiceover_fit"):
                        if field not in segment or segment[field] in (None, "", []):
                            errors.append(f"{segment_prefix}: 缺少字段 {field}")
                    start = segment.get("start")
                    end = segment.get("end")
                    if not isinstance(start, (int, float)) or not isinstance(end, (int, float)) or end <= start:
                        errors.append(f"{segment_prefix}: start/end 无效")
                    elif start < last_end:
                        errors.append(f"{segment_prefix}: 与上一片段重叠或顺序错误")
                    else:
                        last_end = float(end)
    return data, errors


def dedupe(values: list[str]) -> list[str]:
    return list(dict.fromkeys(value for value in values if value))


def infer_segments(track: dict) -> list[dict]:
    """Create useful first-pass sections when manual section annotations are absent."""
    duration = parse_duration(track.get("duration", ""))
    if duration <= 0:
        duration = 180.0
    core = track.get("core_emotions", [])
    fine = track.get("fine_emotions", [])
    atmosphere = track.get("atmosphere", [])
    track_uses = track.get("uses", [])
    if duration <= 30:
        return [{
            "id": "s1",
            "label": "完整短片段",
            "start": 0,
            "end": round(duration, 2),
            "core_emotions": core,
            "fine_emotions": fine,
            "intensity": "中低",
            "atmosphere": atmosphere,
            "uses": dedupe(track_uses + ["故事叙述", "氛围转场", "片尾收束"]),
            "scenes": ["人物近景", "回忆画面", "留白转场"],
            "voiceover_fit": "可用",
            "vocal_presence": "明显人声",
            "annotation": "文件本身为短片段",
        }]

    phase_specs = (
        ("开场铺垫", 0.00, 0.18, "低", ["口播垫底", "知识讲解", "片头铺垫"], ["城市空镜", "产品细节", "环境建立"], "优秀", "轻人声或器乐"),
        ("叙事展开", 0.18, 0.42, "中低", ["故事叙述", "人物故事", "Vlog日常"], ["人物行动", "日常蒙太奇", "采访辅助"], "可用", "可能有人声"),
        ("情绪递进", 0.42, 0.68, "中等", ["情绪递进", "产品展示", "旅行记录"], ["镜头推进", "旅行回忆", "产品演示"], "可用", "明显人声"),
        ("高潮释放", 0.68, 0.88, "高", ["高潮爆发", "快剪卡点", "广告情绪"], ["情绪爆点", "运动快剪", "高光时刻"], "不建议", "明显人声"),
        ("片尾收束", 0.88, 1.00, "中低", ["片尾收束", "氛围转场", "回忆蒙太奇"], ["片尾字幕", "远景离场", "慢镜头回看"], "优秀", "轻人声或器乐"),
    )
    segments = []
    for index, (label, start_ratio, end_ratio, intensity, uses, scenes, fit, vocal) in enumerate(phase_specs, start=1):
        start = round(duration * start_ratio, 2)
        end = round(duration * end_ratio, 2)
        if index in (1, 5):
            extra_uses = [value for value in track_uses if value in {"夜景氛围", "风景空镜", "青春氛围", "氛围转场"}][:1]
        elif index == 2:
            extra_uses = [value for value in track_uses if value in {"人物故事", "感情回忆"}][:1]
        elif index == 3:
            extra_uses = [value for value in track_uses if value in {"情绪递进", "青春氛围"}][:1]
        else:
            extra_uses = [value for value in track_uses if value == "高潮爆发"][:1]
        segments.append({
            "id": f"s{index}",
            "label": label,
            "start": start,
            "end": end,
            "core_emotions": core[:2],
            "fine_emotions": fine[max(0, index - 2):max(0, index - 2) + 3] or fine[:3],
            "intensity": intensity,
            "atmosphere": atmosphere,
            "uses": dedupe(uses + extra_uses),
            "scenes": scenes,
            "voiceover_fit": fit,
            "vocal_presence": vocal,
            "annotation": "自动初标 · 建议试听后微调边界",
        })
    return segments


def build_waveform(audio_path: Path, points: int = 420) -> list[float]:
    ffmpeg = shutil.which("ffmpeg")
    if not ffmpeg:
        return []
    command = [
        ffmpeg,
        "-v", "error",
        "-i", str(audio_path),
        "-ac", "1",
        "-ar", "1000",
        "-f", "s16le",
        "pipe:1",
    ]
    try:
        completed = subprocess.run(command, check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=90)
    except (OSError, subprocess.SubprocessError):
        return []
    samples = array.array("h")
    samples.frombytes(completed.stdout)
    if not samples:
        return []
    stride = max(1, len(samples) // points)
    peaks = []
    for offset in range(0, len(samples), stride):
        window = samples[offset:offset + stride]
        peaks.append(max(abs(value) for value in window) / 32768)
        if len(peaks) == points:
            break
    maximum = max(peaks) or 1
    return [round(value / maximum, 3) for value in peaks]


def enrich_catalog(root: Path, data: dict) -> dict:
    enriched = json.loads(json.dumps(data, ensure_ascii=False))
    for track in enriched["tracks"]:
        if not track.get("segments"):
            track["segments"] = infer_segments(track)
        pure = PurePosixPath(track["relative_path"])
        track["waveform"] = build_waveform(root.joinpath(*pure.parts))
        track["display_name"] = Path(track["file_name"]).stem
    enriched["version"] = max(2, int(enriched.get("version", 1)))
    enriched["segment_taxonomy"] = {
        "uses": ["口播垫底", "知识讲解", "故事叙述", "人物故事", "产品展示", "Vlog日常", "旅行记录", "片头铺垫", "氛围转场", "情绪递进", "高潮爆发", "快剪卡点", "广告情绪", "片尾收束", "回忆蒙太奇"],
        "voiceover_fit": ["优秀", "可用", "不建议"],
        "intensity": ["低", "中低", "中等", "中高", "高"],
    }
    return enriched


def row_values(track: dict) -> list[str]:
    return [
        track["file_name"], track["type"], track["language"],
        "、".join(track["core_emotions"]), "、".join(track["fine_emotions"]), track["intensity"],
        "、".join(track["atmosphere"]), "、".join(track["uses"]),
        track.get("source", "未知"), track.get("license_status", "未知"),
        track.get("duration", "待确认"), track["relative_path"], track.get("notes", ""),
    ]


def segment_rows(tracks: list[dict]) -> list[list[str]]:
    rows = []
    for track in tracks:
        for segment in track["segments"]:
            rows.append([
                track["display_name"], segment["label"], seconds_label(segment["start"]),
                seconds_label(segment["end"]), seconds_label(segment["end"] - segment["start"]),
                "、".join(segment.get("core_emotions", [])), "、".join(segment.get("fine_emotions", [])),
                segment.get("intensity", ""), "、".join(segment.get("uses", [])),
                "、".join(segment.get("scenes", [])), segment.get("voiceover_fit", ""),
                segment.get("vocal_presence", ""), segment.get("annotation", ""),
            ])
    return rows


def build_csv_files(root: Path, tracks: list[dict]) -> list[Path]:
    outputs = []
    for name, headers, rows in (
        ("素材清单.csv", HEADERS, [row_values(track) for track in tracks]),
        ("片段清单.csv", SEGMENT_HEADERS, segment_rows(tracks)),
    ):
        output = root / name
        fd, temp_name = tempfile.mkstemp(prefix=f".{name}.", suffix=".tmp", dir=root)
        try:
            with os.fdopen(fd, "w", encoding="utf-8-sig", newline="") as handle:
                writer = csv.writer(handle)
                writer.writerow(headers)
                writer.writerows(rows)
            os.replace(temp_name, output)
        except Exception:
            try:
                os.unlink(temp_name)
            except FileNotFoundError:
                pass
            raise
        outputs.append(output)
    return outputs


def style_sheet(sheet, widths: tuple[int, ...]) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="135F85")
        cell.alignment = Alignment(horizontal="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def build_xlsx(root: Path, tracks: list[dict]) -> Path | None:
    try:
        from openpyxl import Workbook
    except ImportError:
        return None
    output = root / "素材清单.xlsx"
    workbook = Workbook()
    tracks_sheet = workbook.active
    tracks_sheet.title = "歌曲清单"
    tracks_sheet.append(HEADERS)
    for track in tracks:
        tracks_sheet.append(row_values(track))
    style_sheet(tracks_sheet, (42, 15, 10, 25, 34, 16, 26, 36, 12, 12, 10, 58, 38))

    segments_sheet = workbook.create_sheet("片段清单")
    segments_sheet.append(SEGMENT_HEADERS)
    for row in segment_rows(tracks):
        segments_sheet.append(row)
    style_sheet(segments_sheet, (34, 18, 10, 10, 10, 24, 30, 12, 48, 36, 12, 16, 30))

    help_sheet = workbook.create_sheet("填写说明")
    for row in (
        ("字段", "填写方法"),
        ("核心/细分情绪", "情绪和氛围分开记录；每个片段可覆盖整首歌标签"),
        ("用途", "优先使用口播垫底、知识讲解、故事叙述、产品展示等剪辑目的"),
        ("口播适配", "优秀=不抢人声；可用=降低音量后可用；不建议=人声或高潮明显"),
        ("片段边界", "自动初标只用于快速试听，重要项目请试听后微调"),
    ):
        help_sheet.append(row)
    style_sheet(help_sheet, (22, 90))

    fd, temp_name = tempfile.mkstemp(prefix=".素材清单.", suffix=".xlsx", dir=root)
    os.close(fd)
    try:
        workbook.save(temp_name)
        os.replace(temp_name, output)
    except Exception:
        try:
            os.unlink(temp_name)
        except FileNotFoundError:
            pass
        raise
    return output


def build_html(root: Path, data: dict) -> Path:
    template_path = Path(__file__).resolve().parent.parent / "templates" / "music-library.html"
    if not template_path.is_file():
        raise FileNotFoundError(f"页面模板不存在: {template_path}")
    template = template_path.read_text(encoding="utf-8")
    payload = json.dumps(data, ensure_ascii=False, separators=(",", ":")).replace("<", "\\u003c")
    output = root / "按情绪和用途找歌.html"
    atomic_write_text(output, template.replace("__CATALOG_JSON__", payload))
    return output


def main() -> int:
    args = parse_args()
    root = args.root.expanduser().resolve()
    catalog_path = root / args.catalog
    data, errors = load_and_validate(root, catalog_path)
    if errors:
        print("校验失败:", file=sys.stderr)
        for error in errors:
            print(f"- {error}", file=sys.stderr)
        return 1
    print(f"校验通过: {len(data['tracks'])} 首，音频路径全部存在")
    if args.check_only:
        return 0

    enriched = enrich_catalog(root, data)
    tracks = enriched["tracks"]
    generated = [build_html(root, enriched), *build_csv_files(root, tracks)]
    xlsx = build_xlsx(root, tracks)
    if xlsx:
        generated.append(xlsx)
    else:
        print("提示: 未安装 openpyxl，已生成CSV但跳过XLSX")
    print(f"片段初标: {sum(len(track['segments']) for track in tracks)} 个")
    for path in generated:
        print(f"已生成: {path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
