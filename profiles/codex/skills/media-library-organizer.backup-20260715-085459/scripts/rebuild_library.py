#!/usr/bin/env python3
"""Validate and rebuild the portable audio, video, and image material library."""

from __future__ import annotations

import argparse
import array
import csv
import json
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path, PurePosixPath


AUDIO_REQUIRED = ("file_name", "type", "language", "core_emotions", "fine_emotions", "intensity", "atmosphere", "uses", "relative_path")
AUDIO_ARRAYS = ("core_emotions", "fine_emotions", "atmosphere", "uses")
VIDEO_REQUIRED = ("file_name", "category", "content_tags", "uses", "relative_path", "duration")
VIDEO_ARRAYS = ("content_tags", "uses")
IMAGE_REQUIRED = ("file_name", "category", "content_tags", "uses", "relative_path")
IMAGE_ARRAYS = ("content_tags", "uses")

AUDIO_HEADERS = ("文件名", "类型", "语言", "核心情绪", "细分情绪", "情绪强度", "氛围关键词", "适用场景", "来源", "授权状态", "时长", "相对路径", "备注")
AUDIO_SEGMENT_HEADERS = ("歌曲", "片段名称", "开始", "结束", "片段时长", "核心情绪", "细分情绪", "强度", "用途", "适合画面", "口播适配", "人声状态", "标注来源")
VIDEO_HEADERS = ("文件名", "分类", "内容标签", "用途", "方向", "分辨率", "帧率", "时长", "相对路径", "备注")
VIDEO_SEGMENT_HEADERS = ("视频", "片段名称", "开始", "结束", "片段时长", "内容标签", "景别", "运镜", "用途", "画面质量", "标注来源")
IMAGE_HEADERS = ("文件名", "分类", "内容标签", "用途", "方向", "尺寸", "主色", "相对路径", "备注")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, required=True, help="素材库根目录")
    parser.add_argument("--check-only", action="store_true", help="只校验，不生成文件")
    return parser.parse_args()


def atomic_write_text(path: Path, text: str, encoding: str = "utf-8") -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temp_name = tempfile.mkstemp(prefix=f".{path.name}.", suffix=".tmp", dir=path.parent)
    try:
        with os.fdopen(fd, "w", encoding=encoding, newline="") as handle:
            handle.write(text)
        os.replace(temp_name, path)
    except Exception:
        try:
            os.unlink(temp_name)
        except FileNotFoundError:
            pass
        raise


def parse_duration(value: str | float | int) -> float:
    if isinstance(value, (int, float)):
        return float(value)
    try:
        parts = [float(part) for part in str(value).split(":")]
    except ValueError:
        return 0.0
    if len(parts) == 2:
        return parts[0] * 60 + parts[1]
    if len(parts) == 3:
        return parts[0] * 3600 + parts[1] * 60 + parts[2]
    return 0.0


def seconds_label(seconds: float) -> str:
    seconds = max(0, int(round(seconds)))
    return f"{seconds // 60:02d}:{seconds % 60:02d}"


def dedupe(values: list[str]) -> list[str]:
    return list(dict.fromkeys(value for value in values if isinstance(value, str) and value.strip()))


def load_catalog(root: Path, filename: str, key: str) -> tuple[dict, list[str]]:
    path = root / filename
    if not path.exists():
        return {"version": 1, key: []}, []
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        return {}, [f"{filename}: JSON格式错误: {exc}"]
    if not isinstance(data, dict) or not isinstance(data.get(key), list):
        return data, [f"{filename}: 根对象必须包含 {key} 数组"]
    return data, []


def validate_common(root: Path, items: list[dict], filename: str, required: tuple[str, ...], arrays: tuple[str, ...]) -> list[str]:
    errors: list[str] = []
    seen_names: set[str] = set()
    seen_paths: set[str] = set()
    for index, item in enumerate(items, start=1):
        prefix = f"{filename}[{index}]"
        if not isinstance(item, dict):
            errors.append(f"{prefix}: 必须是对象")
            continue
        for field in required:
            if field not in item or item[field] in (None, "", []):
                errors.append(f"{prefix}: 缺少必填字段 {field}")
        for field in arrays:
            value = item.get(field)
            if not isinstance(value, list) or not value or any(not isinstance(v, str) or not v.strip() for v in value):
                errors.append(f"{prefix}.{field}: 必须是非空字符串数组")
            elif len(value) != len(set(value)):
                errors.append(f"{prefix}.{field}: 包含重复标签")
        name = item.get("file_name")
        rel = item.get("relative_path")
        if isinstance(name, str):
            if name in seen_names:
                errors.append(f"{prefix}: 文件名重复 {name}")
            seen_names.add(name)
        if isinstance(rel, str):
            if "\\" in rel:
                errors.append(f"{prefix}: relative_path 必须使用 / 分隔")
            pure = PurePosixPath(rel)
            if pure.is_absolute() or ".." in pure.parts:
                errors.append(f"{prefix}: relative_path 必须是安全相对路径")
            if rel in seen_paths:
                errors.append(f"{prefix}: 路径重复 {rel}")
            seen_paths.add(rel)
            if isinstance(name, str) and pure.name != name:
                errors.append(f"{prefix}: file_name 与路径末尾不一致")
            if not root.joinpath(*pure.parts).is_file():
                errors.append(f"{prefix}: 素材文件不存在 {rel}")
    return errors


def validate_segments(items: list[dict], filename: str, required: tuple[str, ...]) -> list[str]:
    errors: list[str] = []
    for item_index, item in enumerate(items, start=1):
        segments = item.get("segments")
        if segments is None:
            continue
        if not isinstance(segments, list) or not segments:
            errors.append(f"{filename}[{item_index}].segments: 必须是非空数组")
            continue
        last_end = 0.0
        duration = parse_duration(item.get("duration", 0))
        for segment_index, segment in enumerate(segments, start=1):
            prefix = f"{filename}[{item_index}].segments[{segment_index}]"
            if not isinstance(segment, dict):
                errors.append(f"{prefix}: 必须是对象")
                continue
            for field in required:
                if field not in segment or segment[field] in (None, "", []):
                    errors.append(f"{prefix}: 缺少字段 {field}")
            start = segment.get("start")
            end = segment.get("end")
            if not isinstance(start, (int, float)) or not isinstance(end, (int, float)) or end <= start:
                errors.append(f"{prefix}: start/end 无效")
            elif start < last_end:
                errors.append(f"{prefix}: 与上一片段重叠或顺序错误")
            else:
                last_end = float(end)
                if duration and end > duration + 0.25:
                    errors.append(f"{prefix}: end 超过素材时长")
    return errors


def infer_audio_segments(track: dict) -> list[dict]:
    duration = parse_duration(track.get("duration", "")) or 180.0
    core, fine, atmosphere, track_uses = (track.get("core_emotions", []), track.get("fine_emotions", []), track.get("atmosphere", []), track.get("uses", []))
    if duration <= 30:
        return [{"id": "s1", "label": "完整短片段", "start": 0, "end": round(duration, 2), "core_emotions": core, "fine_emotions": fine, "intensity": "中低", "atmosphere": atmosphere, "uses": dedupe(track_uses + ["故事叙述", "氛围转场", "片尾收束"]), "scenes": ["人物近景", "回忆画面", "留白转场"], "voiceover_fit": "可用", "vocal_presence": "明显人声", "annotation": "文件本身为短片段"}]
    specs = (
        ("开场铺垫", 0.00, 0.18, "低", ["口播垫底", "知识讲解", "片头铺垫"], ["城市空镜", "产品细节", "环境建立"], "优秀", "轻人声或器乐"),
        ("叙事展开", 0.18, 0.42, "中低", ["故事叙述", "人物故事", "Vlog日常"], ["人物行动", "日常蒙太奇", "采访辅助"], "可用", "可能有人声"),
        ("情绪递进", 0.42, 0.68, "中等", ["情绪递进", "产品展示", "旅行记录"], ["镜头推进", "旅行回忆", "产品演示"], "可用", "明显人声"),
        ("高潮释放", 0.68, 0.88, "高", ["高潮爆发", "快剪卡点", "广告情绪"], ["情绪爆点", "运动快剪", "高光时刻"], "不建议", "明显人声"),
        ("片尾收束", 0.88, 1.00, "中低", ["片尾收束", "氛围转场", "回忆蒙太奇"], ["片尾字幕", "远景离场", "慢镜头回看"], "优秀", "轻人声或器乐"),
    )
    segments = []
    for index, (label, a, b, intensity, uses, scenes, fit, vocal) in enumerate(specs, start=1):
        if index in (1, 5):
            extra = [value for value in track_uses if value in {"夜景氛围", "风景空镜", "青春氛围", "氛围转场"}][:1]
        elif index == 2:
            extra = [value for value in track_uses if value in {"人物故事", "感情回忆"}][:1]
        elif index == 3:
            extra = [value for value in track_uses if value in {"情绪递进", "青春氛围"}][:1]
        else:
            extra = [value for value in track_uses if value == "高潮爆发"][:1]
        segments.append({"id": f"s{index}", "label": label, "start": round(duration * a, 2), "end": round(duration * b, 2), "core_emotions": core[:2], "fine_emotions": fine[max(0, index - 2):max(0, index - 2) + 3] or fine[:3], "intensity": intensity, "atmosphere": atmosphere, "uses": dedupe(uses + extra), "scenes": scenes, "voiceover_fit": fit, "vocal_presence": vocal, "annotation": "自动初标 · 建议试听后微调边界"})
    return segments


def infer_video_segments(video: dict) -> list[dict]:
    duration = parse_duration(video.get("duration", "")) or 30.0
    if duration <= 12:
        ratios = (("完整镜头", 0.0, 1.0, "中景", "固定"),)
    else:
        ratios = (("开场镜头", 0.0, 0.22, "全景", "固定"), ("主体展示", 0.22, 0.55, "中景", "轻微运动"), ("细节补充", 0.55, 0.78, "近景", "推进"), ("收束镜头", 0.78, 1.0, "远景", "固定"))
    return [{"id": f"v{index}", "label": label, "start": round(duration * a, 2), "end": round(duration * b, 2), "content_tags": video.get("content_tags", []), "shot_type": shot, "camera_motion": motion, "uses": video.get("uses", []), "quality": "可用", "annotation": "自动初标 · 需逐镜头确认"} for index, (label, a, b, shot, motion) in enumerate(ratios, start=1)]


def build_waveform(audio_path: Path, points: int = 420) -> list[float]:
    ffmpeg = shutil.which("ffmpeg")
    if not ffmpeg:
        return []
    try:
        completed = subprocess.run([ffmpeg, "-v", "error", "-i", str(audio_path), "-ac", "1", "-ar", "1000", "-f", "s16le", "pipe:1"], check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=90)
    except (OSError, subprocess.SubprocessError):
        return []
    samples = array.array("h")
    samples.frombytes(completed.stdout)
    if not samples:
        return []
    stride = max(1, len(samples) // points)
    peaks = [max(abs(value) for value in samples[offset:offset + stride]) / 32768 for offset in range(0, len(samples), stride)][:points]
    maximum = max(peaks) or 1
    return [round(value / maximum, 3) for value in peaks]


def enrich_audio(root: Path, data: dict) -> dict:
    enriched = json.loads(json.dumps(data, ensure_ascii=False))
    for track in enriched["tracks"]:
        track["segments"] = track.get("segments") or infer_audio_segments(track)
        pure = PurePosixPath(track["relative_path"])
        track["waveform"] = build_waveform(root.joinpath(*pure.parts))
        track["display_name"] = Path(track["file_name"]).stem
    enriched["version"] = max(2, int(enriched.get("version", 1)))
    return enriched


def enrich_videos(data: dict) -> dict:
    enriched = json.loads(json.dumps(data, ensure_ascii=False))
    for video in enriched["videos"]:
        video["segments"] = video.get("segments") or infer_video_segments(video)
        video["display_name"] = Path(video["file_name"]).stem
    return enriched


def enrich_images(data: dict) -> dict:
    enriched = json.loads(json.dumps(data, ensure_ascii=False))
    for image in enriched["images"]:
        image["display_name"] = Path(image["file_name"]).stem
    return enriched


def audio_row(track: dict) -> list[str]:
    return [track["file_name"], track["type"], track["language"], "、".join(track["core_emotions"]), "、".join(track["fine_emotions"]), track["intensity"], "、".join(track["atmosphere"]), "、".join(track["uses"]), track.get("source", "未知"), track.get("license_status", "未知"), track.get("duration", "待确认"), track["relative_path"], track.get("notes", "")]


def audio_segment_rows(tracks: list[dict]) -> list[list[str]]:
    return [[track["display_name"], segment["label"], seconds_label(segment["start"]), seconds_label(segment["end"]), seconds_label(segment["end"] - segment["start"]), "、".join(segment.get("core_emotions", [])), "、".join(segment.get("fine_emotions", [])), segment.get("intensity", ""), "、".join(segment.get("uses", [])), "、".join(segment.get("scenes", [])), segment.get("voiceover_fit", ""), segment.get("vocal_presence", ""), segment.get("annotation", "")] for track in tracks for segment in track["segments"]]


def video_row(video: dict) -> list[str]:
    return [video["file_name"], video["category"], "、".join(video["content_tags"]), "、".join(video["uses"]), video.get("orientation", "待确认"), video.get("resolution", "待确认"), str(video.get("fps", "待确认")), video.get("duration", "待确认"), video["relative_path"], video.get("notes", "")]


def video_segment_rows(videos: list[dict]) -> list[list[str]]:
    return [[video["display_name"], segment["label"], seconds_label(segment["start"]), seconds_label(segment["end"]), seconds_label(segment["end"] - segment["start"]), "、".join(segment.get("content_tags", [])), segment.get("shot_type", ""), segment.get("camera_motion", ""), "、".join(segment.get("uses", [])), segment.get("quality", ""), segment.get("annotation", "")] for video in videos for segment in video["segments"]]


def image_row(image: dict) -> list[str]:
    size = image.get("size") or (f"{image.get('width')}×{image.get('height')}" if image.get("width") and image.get("height") else "待确认")
    return [image["file_name"], image["category"], "、".join(image["content_tags"]), "、".join(image["uses"]), image.get("orientation", "待确认"), size, "、".join(image.get("colors", [])), image["relative_path"], image.get("notes", "")]


def write_csv(root: Path, name: str, headers: tuple[str, ...], rows: list[list[str]]) -> Path:
    output = root / name
    fd, temp_name = tempfile.mkstemp(prefix=f".{name}.", suffix=".tmp", dir=root)
    try:
        with os.fdopen(fd, "w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(headers)
            writer.writerows(rows)
        os.replace(temp_name, output)
    except Exception:
        try:
            os.unlink(temp_name)
        except FileNotFoundError:
            pass
        raise
    return output


def build_xlsx(root: Path, bundle: dict) -> Path | None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        return None
    sheets = (
        ("音频清单", AUDIO_HEADERS, [audio_row(track) for track in bundle["audio"]["tracks"]]),
        ("音频片段", AUDIO_SEGMENT_HEADERS, audio_segment_rows(bundle["audio"]["tracks"])),
        ("视频清单", VIDEO_HEADERS, [video_row(video) for video in bundle["video"]["videos"]]),
        ("视频片段", VIDEO_SEGMENT_HEADERS, video_segment_rows(bundle["video"]["videos"])),
        ("图片清单", IMAGE_HEADERS, [image_row(image) for image in bundle["image"]["images"]]),
    )
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, headers, rows in sheets:
        sheet = workbook.create_sheet(title)
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="135F85")
            cell.alignment = Alignment(horizontal="center")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = min(52, max(12, max(len(str(cell.value or "")) for cell in column) + 3))
            for cell in column[1:]:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    output = root / "素材清单.xlsx"
    fd, temp_name = tempfile.mkstemp(prefix=".素材清单.", suffix=".xlsx", dir=root)
    os.close(fd)
    try:
        workbook.save(temp_name)
        os.replace(temp_name, output)
    except Exception:
        try:
            os.unlink(temp_name)
        except FileNotFoundError:
            pass
        raise
    return output


def build_html(root: Path, bundle: dict) -> list[Path]:
    template_path = Path(__file__).resolve().parent.parent / "assets" / "media-library.html"
    if not template_path.is_file():
        raise FileNotFoundError(f"页面模板不存在: {template_path}")
    payload = json.dumps(bundle, ensure_ascii=False, separators=(",", ":")).replace("<", "\\u003c")
    html = template_path.read_text(encoding="utf-8").replace("__CATALOG_JSON__", payload)
    outputs = [root / "素材片段库.html", root / "按情绪和用途找歌.html"]
    for output in outputs:
        atomic_write_text(output, html)
    return outputs


def main() -> int:
    args = parse_args()
    root = args.root.expanduser().resolve()
    audio, audio_errors = load_catalog(root, "歌曲目录.json", "tracks")
    video, video_errors = load_catalog(root, "视频目录.json", "videos")
    image, image_errors = load_catalog(root, "图片目录.json", "images")
    errors = audio_errors + video_errors + image_errors
    if not errors:
        errors += validate_common(root, audio["tracks"], "歌曲目录.json", AUDIO_REQUIRED, AUDIO_ARRAYS)
        errors += validate_segments(audio["tracks"], "歌曲目录.json", ("label", "start", "end", "uses", "voiceover_fit"))
        errors += validate_common(root, video["videos"], "视频目录.json", VIDEO_REQUIRED, VIDEO_ARRAYS)
        errors += validate_segments(video["videos"], "视频目录.json", ("label", "start", "end", "content_tags", "uses"))
        errors += validate_common(root, image["images"], "图片目录.json", IMAGE_REQUIRED, IMAGE_ARRAYS)
    if errors:
        print("校验失败:", file=sys.stderr)
        for error in errors:
            print(f"- {error}", file=sys.stderr)
        return 1
    print(f"校验通过: 音频 {len(audio['tracks'])}，视频 {len(video['videos'])}，图片 {len(image['images'])}")
    if args.check_only:
        return 0

    for filename, key, data in (("视频目录.json", "videos", video), ("图片目录.json", "images", image)):
        path = root / filename
        if not path.exists():
            atomic_write_text(path, json.dumps(data, ensure_ascii=False, indent=2) + "\n")

    bundle = {"version": 1, "audio": enrich_audio(root, audio), "video": enrich_videos(video), "image": enrich_images(image)}
    generated = build_html(root, bundle)
    generated.extend([
        write_csv(root, "素材清单.csv", AUDIO_HEADERS, [audio_row(track) for track in bundle["audio"]["tracks"]]),
        write_csv(root, "片段清单.csv", AUDIO_SEGMENT_HEADERS, audio_segment_rows(bundle["audio"]["tracks"])),
        write_csv(root, "视频片段清单.csv", VIDEO_SEGMENT_HEADERS, video_segment_rows(bundle["video"]["videos"])),
        write_csv(root, "图片清单.csv", IMAGE_HEADERS, [image_row(item) for item in bundle["image"]["images"]]),
    ])
    xlsx = build_xlsx(root, bundle)
    if xlsx:
        generated.append(xlsx)
    print(f"片段索引: 音频 {sum(len(track['segments']) for track in bundle['audio']['tracks'])}，视频 {sum(len(item['segments']) for item in bundle['video']['videos'])}")
    for path in generated:
        print(f"已生成: {path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
